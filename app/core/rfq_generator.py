"""
Moduł do generowania zapytań ofertowych (RFQ - Request for Quotation)
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from typing import Dict, Any


class RFQGenerator:
    """Klasa do generowania dokumentów zapytań ofertowych"""

    def __init__(self):
        self.templates = {
            "basic": self._generate_basic_rfq,
            "detailed": self._generate_detailed_rfq,
            "construction": self._generate_construction_rfq
        }

    def generate(self, data: Dict[str, Any], output_path: str, template: str = "basic"):
        """
        Generuje dokument RFQ

        Args:
            data: Dane do wygenerowania dokumentu
            output_path: Ścieżka do pliku wyjściowego
            template: Typ szablonu (basic, detailed, construction)
        """
        if template not in self.templates:
            template = "basic"

        generator_func = self.templates[template]
        generator_func(data, output_path)

    def _generate_basic_rfq(self, data: Dict[str, Any], output_path: str):
        """Generuje podstawowe zapytanie ofertowe"""
        doc = Document()

        # Tytuł
        title = doc.add_heading('ZAPYTANIE OFERTOWE', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Data
        date_para = doc.add_paragraph()
        date_para.add_run(f"Data: {datetime.now().strftime('%d.%m.%Y')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_paragraph()

        # Informacje o projekcie
        doc.add_heading('1. Informacje podstawowe', 1)

        table = doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'

        cells = table.rows[0].cells
        cells[0].text = 'Nazwa projektu:'
        cells[1].text = data.get('project_name', '')

        cells = table.rows[1].cells
        cells[0].text = 'Firma zamawiająca:'
        cells[1].text = data.get('company_name', '')

        cells = table.rows[2].cells
        cells[0].text = 'Termin składania ofert:'
        cells[1].text = data.get('deadline', '')

        cells = table.rows[3].cells
        cells[0].text = 'Osoba kontaktowa:'
        cells[1].text = data.get('contact_person', 'Do uzupełnienia')

        doc.add_paragraph()

        # Opis projektu
        doc.add_heading('2. Opis przedmiotu zamówienia', 1)
        doc.add_paragraph(data.get('description', ''))

        doc.add_paragraph()

        # Zakres oferty
        doc.add_heading('3. Zakres oferty', 1)
        doc.add_paragraph('Prosimy o przedstawienie oferty obejmującej:')

        items = [
            'Szczegółowy opis proponowanego rozwiązania',
            'Harmonogram realizacji',
            'Wycena (z rozbiciem na poszczególne elementy)',
            'Warunki gwarancji i wsparcia',
            'Kwalifikacje i doświadczenie zespołu'
        ]

        for item in items:
            doc.add_paragraph(item, style='List Bullet')

        doc.add_paragraph()

        # Kryteria oceny
        doc.add_heading('4. Kryteria oceny ofert', 1)
        doc.add_paragraph('Oferty będą oceniane według następujących kryteriów:')

        criteria = [
            'Cena - 40%',
            'Jakość rozwiązania - 30%',
            'Doświadczenie wykonawcy - 20%',
            'Termin realizacji - 10%'
        ]

        for criterion in criteria:
            doc.add_paragraph(criterion, style='List Bullet')

        doc.add_paragraph()

        # Informacje końcowe
        doc.add_heading('5. Sposób składania ofert', 1)
        doc.add_paragraph(
            'Oferty prosimy przesyłać na adres email: '
            f"{data.get('contact_email', 'oferty@firma.pl')} "
            f"do dnia {data.get('deadline', '[termin]')}."
        )

        doc.add_paragraph()
        doc.add_paragraph(
            'Zastrzegamy sobie prawo do odrzucenia ofert bez podania przyczyny '
            'oraz do przeprowadzenia negocjacji z wybranymi oferentami.'
        )

        doc.save(output_path)

    def _generate_detailed_rfq(self, data: Dict[str, Any], output_path: str):
        """Generuje szczegółowe zapytanie ofertowe"""
        doc = Document()

        # Tytuł
        title = doc.add_heading('ZAPYTANIE OFERTOWE - SZCZEGÓŁOWE', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Numer RFQ
        rfq_number = f"RFQ/{datetime.now().strftime('%Y/%m/%d')}/{data.get('project_name', 'XXX')[:3].upper()}"
        para = doc.add_paragraph()
        para.add_run(f"Numer: {rfq_number}").bold = True
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        date_para = doc.add_paragraph()
        date_para.add_run(f"Data: {datetime.now().strftime('%d.%m.%Y')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_page_break()

        # Spis treści (placeholder)
        doc.add_heading('Spis treści', 1)
        toc_items = [
            '1. Informacje podstawowe',
            '2. Opis przedmiotu zamówienia',
            '3. Wymagania techniczne',
            '4. Zakres oferty',
            '5. Warunki realizacji',
            '6. Kryteria oceny',
            '7. Formularz ofertowy'
        ]
        for item in toc_items:
            doc.add_paragraph(item, style='List Number')

        doc.add_page_break()

        # Sekcja 1
        doc.add_heading('1. Informacje podstawowe', 1)

        info_table = doc.add_table(rows=6, cols=2)
        info_table.style = 'Light Grid Accent 1'

        info_data = [
            ('Nazwa projektu:', data.get('project_name', '')),
            ('Zamawiający:', data.get('company_name', '')),
            ('Osoba kontaktowa:', data.get('contact_person', 'Do uzupełnienia')),
            ('Email:', data.get('contact_email', 'oferty@firma.pl')),
            ('Telefon:', data.get('contact_phone', 'Do uzupełnienia')),
            ('Termin składania ofert:', data.get('deadline', ''))
        ]

        for i, (label, value) in enumerate(info_data):
            cells = info_table.rows[i].cells
            cells[0].text = label
            cells[1].text = value

        doc.add_paragraph()

        # Sekcja 2
        doc.add_heading('2. Opis przedmiotu zamówienia', 1)
        doc.add_heading('2.1. Cel projektu', 2)
        doc.add_paragraph(data.get('description', ''))

        doc.add_heading('2.2. Kontekst biznesowy', 2)
        doc.add_paragraph(data.get('business_context', 'Do uzupełnienia'))

        doc.add_heading('2.3. Oczekiwane rezultaty', 2)
        doc.add_paragraph(data.get('expected_results', 'Do uzupełnienia'))

        # Sekcja 3
        doc.add_heading('3. Wymagania techniczne', 1)
        doc.add_paragraph('Szczegółowe wymagania techniczne:')
        doc.add_paragraph('[Do uzupełnienia przez zamawiającego]')

        # Sekcja 4
        doc.add_heading('4. Zakres oferty', 1)
        doc.add_paragraph('Oferta musi zawierać:')

        scope_items = [
            'Opis proponowanego rozwiązania',
            'Harmonogram realizacji z kamieniami milowymi',
            'Szczegółową wycenę z rozbiciem na etapy',
            'Metodologię pracy',
            'Skład zespołu projektowego',
            'Referencje z podobnych projektów',
            'Warunki płatności',
            'Warunki gwarancji',
            'Plan zarządzania ryzykiem'
        ]

        for item in scope_items:
            doc.add_paragraph(item, style='List Bullet')

        # Sekcja 5
        doc.add_heading('5. Warunki realizacji', 1)
        doc.add_paragraph(f"Planowany termin rozpoczęcia: {data.get('start_date', 'Do uzgodnienia')}")
        doc.add_paragraph(f"Planowany termin zakończenia: {data.get('end_date', 'Do uzgodnienia')}")
        doc.add_paragraph(f"Miejsce realizacji: {data.get('location', 'Do uzgodnienia')}")

        # Sekcja 6
        doc.add_heading('6. Kryteria oceny ofert', 1)

        criteria_table = doc.add_table(rows=5, cols=2)
        criteria_table.style = 'Light List Accent 1'

        criteria_table.rows[0].cells[0].text = 'Kryterium'
        criteria_table.rows[0].cells[1].text = 'Waga'

        criteria_data = [
            ('Cena', '40%'),
            ('Jakość rozwiązania technicznego', '25%'),
            ('Doświadczenie i referencje', '20%'),
            ('Harmonogram realizacji', '15%')
        ]

        for i, (criterion, weight) in enumerate(criteria_data, 1):
            criteria_table.rows[i].cells[0].text = criterion
            criteria_table.rows[i].cells[1].text = weight

        doc.add_paragraph()

        # Sekcja 7
        doc.add_heading('7. Formularz ofertowy', 1)
        doc.add_paragraph('Oferent zobowiązany jest do wypełnienia poniższego formularza:')

        form_table = doc.add_table(rows=8, cols=2)
        form_table.style = 'Table Grid'

        form_fields = [
            ('Nazwa firmy:', ''),
            ('NIP:', ''),
            ('Adres:', ''),
            ('Osoba reprezentująca:', ''),
            ('Email kontaktowy:', ''),
            ('Telefon kontaktowy:', ''),
            ('Cena netto:', ''),
            ('Termin realizacji:', '')
        ]

        for i, (label, _) in enumerate(form_fields):
            form_table.rows[i].cells[0].text = label

        doc.add_paragraph()

        # Informacje końcowe
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.add_run('Podpis i pieczęć firmowa').italic = True
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.save(output_path)

    def _generate_construction_rfq(self, data: Dict[str, Any], output_path: str):
        """Generuje zapytanie ofertowe dla branży budowlanej"""
        doc = Document()

        # Tytuł
        title = doc.add_heading('ZAPYTANIE OFERTOWE - ROBOTY BUDOWLANE', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        date_para = doc.add_paragraph()
        date_para.add_run(f"Data: {datetime.now().strftime('%d.%m.%Y')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_paragraph()

        # Informacje podstawowe
        doc.add_heading('1. Informacje o inwestycji', 1)

        table = doc.add_table(rows=5, cols=2)
        table.style = 'Light Grid Accent 1'

        table.rows[0].cells[0].text = 'Nazwa inwestycji:'
        table.rows[0].cells[1].text = data.get('project_name', '')

        table.rows[1].cells[0].text = 'Lokalizacja:'
        table.rows[1].cells[1].text = data.get('location', 'Do uzupełnienia')

        table.rows[2].cells[0].text = 'Inwestor:'
        table.rows[2].cells[1].text = data.get('company_name', '')

        table.rows[3].cells[0].text = 'Termin składania ofert:'
        table.rows[3].cells[1].text = data.get('deadline', '')

        table.rows[4].cells[0].text = 'Planowany termin realizacji:'
        table.rows[4].cells[1].text = data.get('construction_duration', 'Do uzupełnienia')

        doc.add_paragraph()

        # Zakres robót
        doc.add_heading('2. Zakres robót budowlanych', 1)
        doc.add_paragraph(data.get('description', ''))
        doc.add_paragraph()
        doc.add_paragraph('Zakres obejmuje między innymi:')

        construction_items = [
            'Roboty przygotowawcze i rozbiórkowe',
            'Roboty ziemne',
            'Roboty fundamentowe',
            'Roboty konstrukcyjne',
            'Roboty wykończeniowe',
            'Instalacje sanitarne',
            'Instalacje elektryczne'
        ]

        for item in construction_items:
            doc.add_paragraph(item, style='List Bullet')

        doc.add_paragraph()

        # Wymagania
        doc.add_heading('3. Wymagania wobec wykonawcy', 1)
        requirements = [
            'Posiadanie uprawnień budowlanych',
            'Ubezpieczenie OC',
            'Doświadczenie w realizacji podobnych projektów (min. 3 lata)',
            'Referencje z poprzednich realizacji',
            'Możliwość przedstawienia portfolio'
        ]

        for req in requirements:
            doc.add_paragraph(req, style='List Bullet')

        doc.add_paragraph()

        # Dokumentacja
        doc.add_heading('4. Dokumentacja do wglądu', 1)
        doc.add_paragraph('Dokumentacja projektowa dostępna w biurze inwestora po wcześniejszym umówieniu.')
        doc.add_paragraph('Wizja lokalna możliwa po uzgodnieniu terminu.')

        doc.add_paragraph()

        # Składanie ofert
        doc.add_heading('5. Wymagany zakres oferty', 1)
        doc.add_paragraph('Oferta powinna zawierać:')

        offer_items = [
            'Kosztorys ofertowy',
            'Harmonogram rzeczowo-finansowy',
            'Wykaz osób odpowiedzialnych za realizację',
            'Kopia uprawnień budowlanych',
            'Polisa OC',
            'Referencje',
            'Propozycje rozwiązań materiałowych',
            'Warunki płatności',
            'Okres gwarancji'
        ]

        for item in offer_items:
            doc.add_paragraph(item, style='List Bullet')

        doc.add_paragraph()

        # Kryteria oceny
        doc.add_heading('6. Kryteria wyboru wykonawcy', 1)
        criteria = [
            'Cena - 50%',
            'Doświadczenie i referencje - 25%',
            'Termin realizacji - 15%',
            'Jakość proponowanych rozwiązań - 10%'
        ]

        for criterion in criteria:
            doc.add_paragraph(criterion, style='List Bullet')

        doc.add_paragraph()

        # Kontakt
        doc.add_heading('7. Kontakt', 1)
        doc.add_paragraph(f"Osoba kontaktowa: {data.get('contact_person', 'Do uzupełnienia')}")
        doc.add_paragraph(f"Email: {data.get('contact_email', 'oferty@firma.pl')}")
        doc.add_paragraph(f"Telefon: {data.get('contact_phone', 'Do uzupełnienia')}")

        doc.add_paragraph()
        doc.add_paragraph()

        closing = doc.add_paragraph()
        closing.add_run('Z poważaniem,').italic = True
        closing.add_run('\n' + data.get('company_name', '')).bold = True

        doc.save(output_path)
