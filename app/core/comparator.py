"""
Moduł do porównywania ofert Excel
"""

import pandas as pd
from typing import List, Dict, Any
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.dataframe import dataframe_to_rows


class OfferComparator:
    """Klasa do porównywania ofert w formacie Excel"""

    def __init__(self):
        self.red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        self.green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
        self.yellow_fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
        self.bold_font = Font(bold=True)

    def read_offer(self, file_path: str) -> pd.DataFrame:
        """
        Wczytuje ofertę z pliku Excel

        Args:
            file_path: Ścieżka do pliku Excel

        Returns:
            DataFrame z danymi oferty
        """
        try:
            # Próbuj wczytać pierwszy arkusz
            df = pd.read_excel(file_path, sheet_name=0)
            return df
        except Exception as e:
            raise Exception(f"Błąd wczytywania pliku {file_path}: {str(e)}")

    def compare_offers(self, file_paths: List[str]) -> Dict[str, Any]:
        """
        Porównuje wiele ofert Excel

        Args:
            file_paths: Lista ścieżek do plików Excel z ofertami

        Returns:
            Słownik z wynikami porównania
        """
        offers = []
        filenames = []

        # Wczytaj wszystkie oferty
        for path in file_paths:
            df = self.read_offer(path)
            offers.append(df)
            filenames.append(Path(path).name)

        # Podstawowa analiza
        result = {
            "filenames": filenames,
            "count": len(offers),
            "summary": self._generate_summary(offers, filenames),
            "comparison_table": self._create_comparison_table(offers, filenames),
            "price_analysis": self._analyze_prices(offers, filenames),
            "scope_analysis": self._analyze_scope(offers, filenames)
        }

        return result

    def _generate_summary(self, offers: List[pd.DataFrame], filenames: List[str]) -> Dict[str, Any]:
        """Generuje podsumowanie ofert"""
        summary = {
            "total_offers": len(offers),
            "offers_info": []
        }

        for i, (df, filename) in enumerate(zip(offers, filenames)):
            info = {
                "filename": filename,
                "rows": len(df),
                "columns": len(df.columns),
                "column_names": df.columns.tolist()
            }
            summary["offers_info"].append(info)

        return summary

    def _create_comparison_table(self, offers: List[pd.DataFrame], filenames: List[str]) -> List[Dict]:
        """Tworzy tabelę porównawczą"""
        comparison_data = []

        # Znajdź wspólne kolumny
        common_columns = set(offers[0].columns)
        for df in offers[1:]:
            common_columns = common_columns.intersection(set(df.columns))

        # Jeśli są wspólne kolumny, porównaj
        if common_columns:
            for col in sorted(common_columns):
                row = {"parameter": col}
                for i, (df, filename) in enumerate(zip(offers, filenames)):
                    # Weź pierwszą wartość z kolumny (można dostosować)
                    value = df[col].iloc[0] if len(df) > 0 else "N/A"
                    row[f"offer_{i+1}"] = str(value)
                    row[f"offer_{i+1}_name"] = filename
                comparison_data.append(row)

        return comparison_data

    def _analyze_prices(self, offers: List[pd.DataFrame], filenames: List[str]) -> Dict[str, Any]:
        """Analizuje ceny w ofertach"""
        price_columns = []

        # Szukaj kolumn z cenami (zawierające 'cena', 'price', 'kwota', 'wartość')
        for df in offers:
            for col in df.columns:
                col_lower = str(col).lower()
                if any(keyword in col_lower for keyword in ['cena', 'price', 'kwota', 'wartość', 'koszt']):
                    if col not in price_columns:
                        price_columns.append(col)

        prices_data = []
        if price_columns:
            for col in price_columns:
                price_row = {"price_type": col, "values": []}
                for i, (df, filename) in enumerate(zip(offers, filenames)):
                    if col in df.columns:
                        # Spróbuj znaleźć wartości numeryczne
                        values = pd.to_numeric(df[col], errors='coerce').dropna()
                        total = values.sum() if len(values) > 0 else 0
                        price_row["values"].append({
                            "offer": filename,
                            "total": float(total),
                            "count": len(values)
                        })
                prices_data.append(price_row)

            # Znajdź najtańszą ofertę
            if prices_data and prices_data[0]["values"]:
                min_price = min(prices_data[0]["values"], key=lambda x: x["total"])
                max_price = max(prices_data[0]["values"], key=lambda x: x["total"])

                return {
                    "price_columns": price_columns,
                    "prices": prices_data,
                    "cheapest": min_price,
                    "most_expensive": max_price,
                    "difference": max_price["total"] - min_price["total"]
                }

        return {
            "price_columns": price_columns,
            "prices": prices_data,
            "message": "Nie znaleziono kolumn z cenami"
        }

    def _analyze_scope(self, offers: List[pd.DataFrame], filenames: List[str]) -> Dict[str, Any]:
        """Analizuje zakres prac w ofertach"""
        scope_data = {
            "items_count": [],
            "unique_items": []
        }

        for i, (df, filename) in enumerate(zip(offers, filenames)):
            scope_data["items_count"].append({
                "offer": filename,
                "total_items": len(df)
            })

        return scope_data

    def save_comparison(self, result: Dict[str, Any], output_path: str):
        """
        Zapisuje wyniki porównania do pliku Excel z formatowaniem

        Args:
            result: Wyniki porównania
            output_path: Ścieżka do pliku wyjściowego
        """
        wb = openpyxl.Workbook()

        # Arkusz 1: Podsumowanie
        ws_summary = wb.active
        ws_summary.title = "Podsumowanie"

        ws_summary['A1'] = "Porównanie Ofert"
        ws_summary['A1'].font = self.bold_font
        ws_summary['A1'].fill = self.yellow_fill

        row = 3
        ws_summary[f'A{row}'] = "Liczba porównanych ofert:"
        ws_summary[f'B{row}'] = result['count']

        row += 2
        ws_summary[f'A{row}'] = "Pliki:"
        ws_summary[f'A{row}'].font = self.bold_font
        row += 1

        for filename in result['filenames']:
            ws_summary[f'A{row}'] = filename
            row += 1

        # Arkusz 2: Tabela porównawcza
        if result.get('comparison_table'):
            ws_comparison = wb.create_sheet("Porównanie")

            # Nagłówki
            headers = ["Parameter"]
            for i in range(result['count']):
                headers.append(f"Oferta {i+1}")

            for col, header in enumerate(headers, 1):
                cell = ws_comparison.cell(1, col, header)
                cell.font = self.bold_font
                cell.fill = self.yellow_fill

            # Dane
            for row_idx, row_data in enumerate(result['comparison_table'], 2):
                ws_comparison.cell(row_idx, 1, row_data['parameter'])
                for i in range(result['count']):
                    value = row_data.get(f'offer_{i+1}', 'N/A')
                    ws_comparison.cell(row_idx, i+2, value)

        # Arkusz 3: Analiza cen
        if result.get('price_analysis') and result['price_analysis'].get('prices'):
            ws_prices = wb.create_sheet("Analiza Cen")

            ws_prices['A1'] = "Analiza Cen"
            ws_prices['A1'].font = self.bold_font
            ws_prices['A1'].fill = self.yellow_fill

            row = 3
            price_analysis = result['price_analysis']

            if 'cheapest' in price_analysis:
                ws_prices[f'A{row}'] = "Najtańsza oferta:"
                ws_prices[f'B{row}'] = price_analysis['cheapest']['offer']
                ws_prices[f'C{row}'] = price_analysis['cheapest']['total']
                ws_prices[f'B{row}'].fill = self.green_fill
                row += 1

                ws_prices[f'A{row}'] = "Najdroższa oferta:"
                ws_prices[f'B{row}'] = price_analysis['most_expensive']['offer']
                ws_prices[f'C{row}'] = price_analysis['most_expensive']['total']
                ws_prices[f'B{row}'].fill = self.red_fill
                row += 1

                ws_prices[f'A{row}'] = "Różnica:"
                ws_prices[f'C{row}'] = price_analysis['difference']
                ws_prices[f'C{row}'].font = self.bold_font

        wb.save(output_path)
