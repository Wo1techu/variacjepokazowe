"""
Główna aplikacja FastAPI - OfferComparator
"""

from fastapi import FastAPI, UploadFile, File, Request, Form
from fastapi.responses import HTMLResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from typing import List
import os
from pathlib import Path

from app.core.comparator import OfferComparator
from app.core.rfq_generator import RFQGenerator
from app.core.file_extractor import FileExtractor, OfferDataParser
from app.core.claude_analyzer import ClaudeOfferAnalyzer

# Inicjalizacja FastAPI
app = FastAPI(
    title="OfferComparator",
    description="Narzędzie do porównywania ofert Excel i generowania zapytań ofertowych",
    version="1.0.0"
)

# Ścieżki
BASE_DIR = Path(__file__).resolve().parent.parent
UPLOAD_DIR = BASE_DIR / "app" / "uploads"
OUTPUT_DIR = BASE_DIR / "app" / "outputs"
STATIC_DIR = BASE_DIR / "app" / "static"
TEMPLATES_DIR = BASE_DIR / "app" / "templates"

# Upewnij się że katalogi istnieją
UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

# Montowanie plików statycznych i szablonów
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
templates = Jinja2Templates(directory=str(TEMPLATES_DIR))

# Inicjalizacja komponentów
comparator = OfferComparator()
rfq_generator = RFQGenerator()
file_extractor = FileExtractor()
data_parser = OfferDataParser()

# Inicjalizacja Claude AI (z fallbackiem na regex parser)
claude_analyzer = None
use_claude = os.getenv("USE_CLAUDE_AI", "true").lower() == "true"

if use_claude:
    try:
        claude_analyzer = ClaudeOfferAnalyzer()
        print("✅ Claude AI analyzer initialized successfully")
    except ValueError as e:
        print(f"⚠️  Claude AI not available: {e}")
        print("📝 Using fallback regex parser instead")
    except Exception as e:
        print(f"⚠️  Error initializing Claude AI: {e}")
        print("📝 Using fallback regex parser instead")


@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    """Strona główna aplikacji"""
    return templates.TemplateResponse("index.html", {"request": request})


@app.get("/compare", response_class=HTMLResponse)
async def compare_page(request: Request):
    """Strona porównywania ofert"""
    return templates.TemplateResponse("compare.html", {"request": request})


@app.get("/generate-rfq", response_class=HTMLResponse)
async def rfq_page(request: Request):
    """Strona generowania zapytań ofertowych"""
    return templates.TemplateResponse("rfq.html", {"request": request})


@app.post("/api/compare")
async def compare_offers(files: List[UploadFile] = File(...), plain_text: str = Form(None)):
    """
    Endpoint do porównywania przesłanych ofert (PDF, Word, JPG, Excel, tekst)
    """
    if not files and not plain_text:
        return {"error": "Prześlij co najmniej 2 pliki lub wklej tekst"}

    if files and len(files) < 2 and not plain_text:
        return {"error": "Potrzebujesz co najmniej 2 plików do porównania"}

    # Zapisz przesłane pliki
    saved_files = []
    extracted_offers = []

    for file in files:
        file_path = UPLOAD_DIR / file.filename
        with open(file_path, "wb") as f:
            content = await file.read()
            f.write(content)
        saved_files.append(str(file_path))

    try:
        # Ekstraktuj dane z każdego pliku
        for file_path in saved_files:
            try:
                # Wyciągnij tekst z pliku
                text = file_extractor.extract_text(file_path)

                # Parsuj tekst - użyj Claude AI jeśli dostępny, w przeciwnym razie regex
                if claude_analyzer:
                    offer_data = claude_analyzer.analyze_offer(text, Path(file_path).name)
                else:
                    offer_data = data_parser.parse(text, Path(file_path).name)

                extracted_offers.append(offer_data)

            except Exception as e:
                # Jeśli nie udało się wyekstrahować, dodaj informację o błędzie
                extracted_offers.append({
                    "source_file": Path(file_path).name,
                    "error": str(e),
                    "extracted_data": {}
                })

        # Jeśli był tekst wklejony, również go przeanalizuj
        if plain_text:
            if claude_analyzer:
                offer_data = claude_analyzer.analyze_offer(plain_text, "Wklejony tekst")
            else:
                offer_data = data_parser.parse(plain_text, "Wklejony tekst")
            extracted_offers.append(offer_data)

        # Przygotuj dane do porównania
        result = {
            "count": len(extracted_offers),
            "offers": extracted_offers,
            "comparison": _create_comparison_summary(extracted_offers)
        }

        # Zapisz wynik do pliku Excel
        output_filename = f"porownanie_{len(extracted_offers)}_ofert.xlsx"
        output_path = OUTPUT_DIR / output_filename
        _save_comparison_to_excel(extracted_offers, str(output_path))

        return {
            "success": True,
            "data": result,
            "download_url": f"/download/{output_filename}"
        }
    except Exception as e:
        import traceback
        return {"error": f"Błąd podczas porównywania: {str(e)}", "trace": traceback.format_exc()}
    finally:
        # Wyczyść przesłane pliki
        for file_path in saved_files:
            if os.path.exists(file_path):
                os.remove(file_path)


def _create_comparison_summary(offers: List[Dict]) -> Dict:
    """Tworzy podsumowanie porównania ofert"""
    prices = []
    companies = []

    for offer in offers:
        if "extracted_data" in offer:
            data = offer["extracted_data"]
            if "total_price" in data:
                prices.append({
                    "source": offer.get("source_file", "Unknown"),
                    "price": data["total_price"]
                })
            if "company_name" in data:
                companies.append(data["company_name"])

    summary = {
        "total_offers": len(offers),
        "offers_with_prices": len(prices)
    }

    if prices:
        sorted_prices = sorted(prices, key=lambda x: x["price"])
        summary["cheapest"] = sorted_prices[0]
        summary["most_expensive"] = sorted_prices[-1]
        summary["price_difference"] = sorted_prices[-1]["price"] - sorted_prices[0]["price"]

    if companies:
        summary["companies"] = companies

    return summary


def _save_comparison_to_excel(offers: List[Dict], output_path: str):
    """Zapisuje porównanie do pliku Excel"""
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font

    # Przygotuj dane do DataFrame
    rows = []
    for offer in offers:
        row = {
            "Źródło": offer.get("source_file", "N/A"),
        }

        if "extracted_data" in offer:
            data = offer["extracted_data"]
            row["Firma"] = data.get("company_name", "N/A")
            row["NIP"] = data.get("nip", "N/A")
            row["Cena całkowita"] = data.get("total_price", "N/A")
            row["Email"] = ", ".join(data.get("emails", [])) if data.get("emails") else "N/A"
            row["Telefon"] = ", ".join(data.get("phones", [])) if data.get("phones") else "N/A"
            row["Daty"] = ", ".join(data.get("dates", [])) if data.get("dates") else "N/A"

        if "error" in offer:
            row["Błąd"] = offer["error"]

        rows.append(row)

    # Utwórz DataFrame i zapisz
    df = pd.DataFrame(rows)
    df.to_excel(output_path, index=False, sheet_name="Porównanie")

    # Dodaj formatowanie
    wb = load_workbook(output_path)
    ws = wb.active

    # Nagłówki - żółte tło
    yellow_fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
    bold_font = Font(bold=True)

    for cell in ws[1]:
        cell.fill = yellow_fill
        cell.font = bold_font

    # Znajdź najniższą i najwyższą cenę i zakoloruj
    price_col = None
    for idx, cell in enumerate(ws[1], 1):
        if "Cena" in str(cell.value):
            price_col = idx
            break

    if price_col:
        prices = []
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, price_col)
            if cell.value and cell.value != "N/A":
                try:
                    prices.append((row_idx, float(cell.value)))
                except (ValueError, TypeError):
                    pass

        if prices:
            prices.sort(key=lambda x: x[1])
            green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
            red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")

            # Najniższa cena - zielone
            ws.cell(prices[0][0], price_col).fill = green_fill
            # Najwyższa cena - czerwone
            ws.cell(prices[-1][0], price_col).fill = red_fill

    wb.save(output_path)


@app.post("/api/generate-rfq")
async def generate_rfq(
    project_name: str = Form(...),
    company_name: str = Form(...),
    deadline: str = Form(...),
    description: str = Form(...),
    template: str = Form("basic")
):
    """
    Endpoint do generowania zapytania ofertowego
    """
    try:
        data = {
            "project_name": project_name,
            "company_name": company_name,
            "deadline": deadline,
            "description": description
        }

        # Generuj dokument
        output_filename = f"zapytanie_ofertowe_{project_name.replace(' ', '_')}.docx"
        output_path = OUTPUT_DIR / output_filename

        rfq_generator.generate(data, str(output_path), template)

        return {
            "success": True,
            "message": "Zapytanie ofertowe wygenerowane pomyślnie",
            "download_url": f"/download/{output_filename}"
        }
    except Exception as e:
        return {"error": f"Błąd podczas generowania: {str(e)}"}


@app.get("/download/{filename}")
async def download_file(filename: str):
    """
    Pobieranie wygenerowanych plików
    """
    file_path = OUTPUT_DIR / filename
    if file_path.exists():
        return FileResponse(
            path=str(file_path),
            filename=filename,
            media_type='application/octet-stream'
        )
    return {"error": "Plik nie istnieje"}


@app.get("/health")
async def health_check():
    """Health check endpoint"""
    return {"status": "ok", "version": "1.0.0"}
